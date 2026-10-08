# ------------------------------------------------------------------------------------------
# yiding_file_io
# ------------------------------------------------------------------------------------------

from __future__ import annotations

import csv
import hashlib
import json
import os
from collections.abc import Mapping
from copy import deepcopy
from datetime import datetime
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path
from typing import Any
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENTATION


FORMAT_NAME = "yiding"
FORMAT_VERSION = 1
YIDING_EXTENSION = ".yiding"

PROJECT_SETTING_KEYS = (
    "punctuation",
    "translation",
    "term_extraction",
    "glossary_extraction",
)

PUNCTUATION_KEYS = (
    "guidelines",
    "facsimile_span",
    "punctuation_span",
    "max_unsegmented_span",
    "max_punctuation_attempts",
    "punctuation",
    "cross_check",
    "cross_examination",
    "cross_correction",
)

TRANSLATION_KEYS = (
    "guidelines",
    "language",
    "glossary",
    "translation_span",
    "translation",
    "glossary_extraction",
    "cross_check",
    "cross_examination",
    "cross_correction",
    "llm_glossary_selection",
    "glossary_selection",
)

EXTRACTION_KEYS = (
    "guidelines",
    "span",
    "extraction",
)


class YidingFileError(Exception):
    """Raised when a .yiding project cannot be read, written, or validated."""


# ------------------------------------------------------------------------------------------
# BASIC HELPERS
# ------------------------------------------------------------------------------------------

def _now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def _get_yiding_version() -> str:
    try:
        return version("yiding")
    except PackageNotFoundError:
        return "0.1.0"


def _plain_json(value: Any) -> Any:
    """Convert TOMLKit/Python containers to ordinary JSON-compatible values."""
    unwrap = getattr(value, "unwrap", None)

    if callable(unwrap):
        try:
            value = unwrap()
        except Exception:
            pass

    if isinstance(value, Mapping):
        return {
            str(key): _plain_json(item)
            for key, item in value.items()
        }

    if isinstance(value, (list, tuple)):
        return [_plain_json(item) for item in value]

    if value is None or isinstance(value, (str, int, float, bool)):
        return value

    raise TypeError(
        f"Value of type {type(value).__name__} is not JSON serializable: {value!r}"
    )


def _copy_selected(mapping: Mapping[str, Any], keys: tuple[str, ...]) -> dict[str, Any]:
    return {
        key: _plain_json(mapping[key])
        for key in keys
        if key in mapping
    }


def ensure_yiding_extension(path: str | Path) -> Path:
    path = Path(path)

    if path.suffix.lower() != YIDING_EXTENSION:
        path = path.with_suffix(YIDING_EXTENSION)

    return path


def get_schema_path() -> Path:
    return Path(__file__).parent / "schemas" / "yiding.schema.json"


# ------------------------------------------------------------------------------------------
# SETTINGS
# ------------------------------------------------------------------------------------------

def project_settings_to_dict(settings: Mapping[str, Any]) -> dict[str, Any]:
    """
    Copy only project settings into the .yiding file.

    API settings and API secrets are intentionally excluded.
    """
    settings = _plain_json(settings)

    missing = [
        key
        for key in PROJECT_SETTING_KEYS
        if key not in settings
    ]

    if missing:
        raise YidingFileError(
            "Missing project settings section(s): " + ", ".join(missing)
        )

    return {
        "punctuation": _copy_selected(
            settings["punctuation"],
            PUNCTUATION_KEYS,
        ),
        "translation": _copy_selected(
            settings["translation"],
            TRANSLATION_KEYS,
        ),
        "term_extraction": _copy_selected(
            settings["term_extraction"],
            EXTRACTION_KEYS,
        ),
        "glossary_extraction": _copy_selected(
            settings["glossary_extraction"],
            EXTRACTION_KEYS,
        ),
    }


def snapshot_operation_settings(
    project_settings: Mapping[str, Any],
    operation: str,
) -> dict[str, Any]:
    if operation not in PROJECT_SETTING_KEYS:
        raise YidingFileError(f"Unknown operation: {operation}")

    return deepcopy(_plain_json(project_settings[operation]))


# ------------------------------------------------------------------------------------------
# SOURCE
# ------------------------------------------------------------------------------------------

def _text_type(text: Any) -> str:
    if hasattr(text, "kanripo_code"):
        return "kanripo"

    return "raw"


def source_to_dict(
    text: Any,
    source_override: Mapping[str, Any] | None = None,
) -> dict[str, Any]:

    if source_override is not None:
        return _plain_json(source_override)

    if _text_type(text) == "kanripo":
        return {
            "kanripo_code": str(
                getattr(text, "kanripo_code")
            ),
            "glosses_on": bool(
                getattr(text, "glosses_on", True)
            ),
        }

    required = (
        "unpunctuated_divider",
        "unpunctuated_keep_divider",
        "unpunctuated_divider_is_whole_line",

        "punctuated_divider",
        "punctuated_keep_divider",
        "punctuated_divider_is_whole_line",

        "translation_divider",
        "translation_keep_divider",
        "translation_divider_is_whole_line",
    )

    missing = [
        name
        for name in required
        if not hasattr(text, name)
    ]

    if missing:
        raise YidingFileError(
            "RawText does not preserve its segmentation "
            "source settings: "
            + ", ".join(missing)
        )

    return {
        "unpunctuated": {
            "divider": text.unpunctuated_divider,
            "keep_divider": bool(
                text.unpunctuated_keep_divider
            ),
            "divider_is_whole_line": bool(
                text.unpunctuated_divider_is_whole_line
            ),
        },

        "punctuated": {
            "divider": text.punctuated_divider,
            "keep_divider": bool(
                text.punctuated_keep_divider
            ),
            "divider_is_whole_line": bool(
                text.punctuated_divider_is_whole_line
            ),
        },

        "translation": {
            "divider": text.translation_divider,
            "keep_divider": bool(
                text.translation_keep_divider
            ),
            "divider_is_whole_line": bool(
                text.translation_divider_is_whole_line
            ),
        },
    }


# ------------------------------------------------------------------------------------------
# GLOSSARIES
# ------------------------------------------------------------------------------------------

def hash_file(path: str | Path) -> str:
    path = Path(path)
    digest = hashlib.sha256()

    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)

    return digest.hexdigest()


def count_csv_entries(path: str | Path) -> int:
    path = Path(path)
    count = 0

    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)

        for row in reader:
            if row and any(cell.strip() for cell in row):
                count += 1

    return count


def glossary_source_info(path: str | Path | None) -> dict[str, Any]:
    """
    Record the original user glossary without embedding its contents.
    """
    if not path:
        return {
            "filename": None,
            "sha256": None,
            "entry_count": None,
        }

    path = Path(path)

    if not path.exists():
        return {
            "filename": path.name,
            "sha256": None,
            "entry_count": None,
        }

    return {
        "filename": path.name,
        "sha256": hash_file(path),
        "entry_count": count_csv_entries(path),
    }


def glossary_dict_to_rows(glossary: Mapping[str, Any] | None) -> list[dict[str, Any]]:
    """
    Convert a runtime Yiding glossary dictionary into JSON rows.

    Canonical runtime shape:
        {"道": ["dào", {"Way", "path"}]}

    Also accepts:
        {"道": ["dào", "Way", "path"]}

    and:
        {"道": {"pinyin": "dào", "translations": ["Way", "path"]}}
    """
    if not glossary:
        return []

    rows = []

    for chinese, value in glossary.items():
        if isinstance(value, Mapping):
            pinyin = str(value.get("pinyin", ""))
            translations = value.get("translations", [])

            if isinstance(translations, str):
                translations = [translations]

            translations = [str(item) for item in translations]

        elif isinstance(value, (list, tuple)):
            values = list(value)
            pinyin = str(values[0]) if values else ""
            
            if (len(values) > 1 and isinstance(values[1], (set, list, tuple))):
                translations = [str(item) for item in values[1]]
            else:
                translations = [str(item) for item in values[1:]]
                
        elif value is None:
            pinyin = ""
            translations = []

        else:
            pinyin = ""
            translations = [str(value)]

        rows.append({
            "chinese": str(chinese),
            "pinyin": pinyin,
            "translations": translations,
        })

    return rows


def glossary_rows_to_dict(
    rows: list[Mapping[str, Any]],
) -> dict[str, list[Any]]:

    glossary = {}

    for row in rows:
        chinese = str(row["chinese"])
        pinyin = str(row.get("pinyin", ""))
        translations = {str(item) for item in row.get("translations", [])}
        glossary[chinese] = [pinyin, translations,]

    return glossary


# ------------------------------------------------------------------------------------------
# TEXT DATA
# ------------------------------------------------------------------------------------------

def metadata_to_dict(text: Any) -> dict[str, Any]:
    return {
        "original_title": getattr(text, "original_title", None),
        "translated_title": getattr(text, "translated_title", None),
        "translation_language": getattr(text, "translation_language",None,),
        "properties": _plain_json(getattr(text, "properties", [])),
        "is_structured": bool(getattr(text, "is_structured", False)),
        "is_punctuated": bool(getattr(text, "is_punctuated", False)),
        "is_translated": bool(getattr(text, "is_translated", False)),
        "is_term_extracted": bool(getattr(text, "is_term_extracted", False)),
        "is_glossary_extracted": bool(getattr(text, "is_glossary_extracted", False)),
    }


def text_data_to_dict(text: Any) -> dict[str, Any]:
    return {
        "segments": _plain_json(getattr(text, "segments", [])),
        "unpunctuated_segments": _plain_json(
            getattr(text, "unpunctuated_segments", [])
        ),
        "punctuated_segments": _plain_json(
            getattr(text, "punctuated_segments", [])
        ),
        "translated_segments": _plain_json(
            getattr(text, "translated_segments", [])
        ),
        "page_labels": _plain_json(getattr(text, "page_labels", [])),
        "page_lines": _plain_json(getattr(text, "page_lines", [])),
        "segment_labels": _plain_json(getattr(text, "segment_labels", [])),
        "glosses": _plain_json(getattr(text, "glosses", [])),
        "punctuated_glosses": _plain_json(
            getattr(text, "punctuated_glosses", [])
        ),
    }


def default_provenance(text: Any) -> dict[str, list[str | None]]:
    return {
        "punctuated_segment_run_ids": [
            None
            for _ in getattr(text, "punctuated_segments", [])
        ],
        "translated_segment_run_ids": [
            None
            for _ in getattr(text, "translated_segments", [])
        ],
    }


def workflow_to_dict(
    text: Any,
    workflow_override: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    workflow = {
        "status": "idle",
        "operation": None,
        "active_run_id": None,
        "working_i": int(getattr(text, "working_i", 0)),
        "working_text": str(getattr(text, "working_text", "")),
        "working_labels": _plain_json(
            getattr(text, "working_labels", [])
        ),
        "translation_i": int(getattr(text, "translation_i", 0)),
        "term_extraction_i": int(getattr(text, "term_extraction_i", 0)),
        "glossary_extraction_i": int(getattr(text, "glossary_extraction_i", 0)),
        "autosaved_at": None,
    }

    if workflow_override:
        workflow.update(_plain_json(workflow_override))

    return workflow


# ------------------------------------------------------------------------------------------
# RUN / PROVENANCE HELPERS
# ------------------------------------------------------------------------------------------

def create_run(
    operation: str,
    project_settings: Mapping[str, Any],
    run_id: str,
) -> dict[str, Any]:
    return {
        "id": run_id,
        "operation": operation,
        "status": "in_progress",
        "started_at": _now_iso(),
        "completed_at": None,
        "settings": snapshot_operation_settings(
            project_settings,
            operation,
        ),
    }


def finish_run(
    run: dict[str, Any],
    status: str = "completed",
) -> None:
    allowed = {
        "completed",
        "failed",
        "cancelled",
    }

    if status not in allowed:
        raise YidingFileError(
            f"Invalid final run status: {status}"
        )

    run["status"] = status
    run["completed_at"] = _now_iso()


# ------------------------------------------------------------------------------------------
# PROJECT BUILDING
# ------------------------------------------------------------------------------------------

def build_project_data(
    text: Any,
    settings: Mapping[str, Any] | None = None,
    *,
    source_override: Mapping[str, Any] | None = None,
    workflow: Mapping[str, Any] | None = None,
    runs: list[Mapping[str, Any]] | None = None,
    provenance: Mapping[str, Any] | None = None,
    glossary_source_path: str | Path | None = None,
    created_at: str | None = None,
) -> dict[str, Any]:
    if settings is None:
        settings = getattr(text, "settings", None)

    if settings is None:
        raise YidingFileError(
            "No project settings were supplied and text.settings is None."
        )

    project_settings = project_settings_to_dict(settings)

    if glossary_source_path is None:
        glossary_source_path = (
            project_settings
            .get("translation", {})
            .get("glossary")
        )

    translation_glossary = getattr(
        text,
        "translation_glossary",
        {},
    )

    new_glossary = getattr(
        text,
        "new_glossary",
        {},
    )

    extracted_terms = getattr(
        text,
        "extracted_terms",
        {},
    )

    extracted_glossary = getattr(
        text,
        "extracted_glossary",
        {},
    )
    
    return {
        "format": FORMAT_NAME,
        "format_version": FORMAT_VERSION,
        "yiding_version": _get_yiding_version(),
        "created_at": created_at or _now_iso(),
        "modified_at": _now_iso(),
        "text_type": _text_type(text),
        "metadata": metadata_to_dict(text),
        "source": source_to_dict(
            text,
            source_override=source_override,
        ),
        "settings": project_settings,
        "workflow": workflow_to_dict(
            text,
            workflow_override=workflow,
        ),
        "runs": _plain_json(runs or []),
        "provenance": _plain_json(
            provenance
            if provenance is not None
            else default_provenance(text)
        ),
        "text": text_data_to_dict(text),
        "glossaries": {
            "source": glossary_source_info(
                glossary_source_path),
            "translation_glossary": {
                "rows": glossary_dict_to_rows(
                    translation_glossary
                )
            },

            "new_glossary": {
                "rows": glossary_dict_to_rows(
                    new_glossary
                )
            },

            "extracted_terms": {
                "rows": glossary_dict_to_rows(
                    extracted_terms
                )
            },
            "extracted_glossary": {
                "rows":
                    glossary_dict_to_rows(
                        extracted_glossary
                    )
            },            
        },
        "log": str(getattr(text, "log", "")),
    }


# ------------------------------------------------------------------------------------------
# SCHEMA VALIDATION
# ------------------------------------------------------------------------------------------

def load_schema() -> dict[str, Any]:
    path = get_schema_path()

    if not path.exists():
        raise YidingFileError(
            f"Yiding JSON schema was not found: {path}"
        )

    try:
        with path.open("r", encoding="utf-8") as file:
            return json.load(file)

    except json.JSONDecodeError as exc:
        raise YidingFileError(
            f"Invalid yiding.schema.json: {exc}"
        ) from exc


def validate_project_data(project: Mapping[str, Any]) -> None:
    """Validate project data against schemas/yiding.schema.json."""
    try:
        import jsonschema

    except ImportError as exc:
        raise YidingFileError(
            "JSON Schema validation requires the 'jsonschema' package. "
            "Add jsonschema to Yiding's dependencies."
        ) from exc

    schema = load_schema()

    try:
        jsonschema.validate(
            instance=_plain_json(project),
            schema=schema,
        )

    except jsonschema.ValidationError as exc:
        location = ".".join(
            str(part)
            for part in exc.absolute_path
        )

        if location:
            location = f" at '{location}'"

        raise YidingFileError(
            f"Invalid Yiding project{location}: {exc.message}"
        ) from exc


# ------------------------------------------------------------------------------------------
# SAVE / AUTOSAVE
# ------------------------------------------------------------------------------------------

def _existing_created_at(path: Path) -> str | None:
    if not path.exists():
        return None

    try:
        with path.open("r", encoding="utf-8") as file:
            existing = json.load(file)

        created_at = existing.get("created_at")

        if isinstance(created_at, str):
            return created_at

    except (OSError, json.JSONDecodeError):
        pass

    return None


def write_project_data(
    path: str | Path,
    project: Mapping[str, Any],
    *,
    validate: bool = True,
) -> Path:
    """
    Atomically write a .yiding file.

    The temporary-file + os.replace approach helps prevent an autosave crash
    from leaving a half-written project file.
    """
    path = ensure_yiding_extension(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    project = deepcopy(_plain_json(project))
    project["modified_at"] = _now_iso()

    if validate:
        validate_project_data(project)

    temp_path = path.with_name(
        f".{path.name}.tmp"
    )

    try:
        with temp_path.open(
            "w",
            encoding="utf-8",
            newline="\n",
        ) as file:
            json.dump(
                project,
                file,
                ensure_ascii=False,
                indent=2,
            )

            file.write("\n")
            file.flush()
            os.fsync(file.fileno())

        os.replace(
            temp_path,
            path,
        )

    except Exception:
        try:
            temp_path.unlink(missing_ok=True)
        except OSError:
            pass

        raise

    return path


def save_project(
    path: str | Path,
    text: Any,
    settings: Mapping[str, Any] | None = None,
    *,
    source_override: Mapping[str, Any] | None = None,
    workflow: Mapping[str, Any] | None = None,
    runs: list[Mapping[str, Any]] | None = None,
    provenance: Mapping[str, Any] | None = None,
    glossary_source_path: str | Path | None = None,
    validate: bool = True,
) -> Path:
    path = ensure_yiding_extension(path)

    project = build_project_data(
        text,
        settings,
        source_override=source_override,
        workflow=workflow,
        runs=runs,
        provenance=provenance,
        glossary_source_path=glossary_source_path,
        created_at=_existing_created_at(path),
    )

    return write_project_data(
        path,
        project,
        validate=validate,
    )


def autosave_project(
    path: str | Path,
    text: Any,
    settings: Mapping[str, Any] | None = None,
    *,
    workflow: Mapping[str, Any],
    runs: list[Mapping[str, Any]] | None = None,
    provenance: Mapping[str, Any] | None = None,
    glossary_source_path: str | Path | None = None,
    source_override: Mapping[str, Any] | None = None,
    validate: bool = True,
) -> Path:
    workflow = dict(_plain_json(workflow))
    workflow["autosaved_at"] = _now_iso()

    return save_project(
        path,
        text,
        settings,
        source_override=source_override,
        workflow=workflow,
        runs=runs,
        provenance=provenance,
        glossary_source_path=glossary_source_path,
        validate=validate,
    )


# ------------------------------------------------------------------------------------------
# LOAD
# ------------------------------------------------------------------------------------------

def load_project(
    path: str | Path,
    *,
    validate: bool = True,
) -> dict[str, Any]:
    path = Path(path)

    if not path.exists():
        raise YidingFileError(
            f"Yiding project does not exist: {path}"
        )

    try:
        with path.open("r", encoding="utf-8") as file:
            project = json.load(file)

    except json.JSONDecodeError as exc:
        raise YidingFileError(
            f"Invalid JSON in {path.name}: {exc}"
        ) from exc

    if project.get("format") != FORMAT_NAME:
        raise YidingFileError(
            f"{path.name} is not a Yiding project file."
        )

    if project.get("format_version") != FORMAT_VERSION:
        raise YidingFileError(
            "Unsupported Yiding format version: "
            f"{project.get('format_version')!r}. "
            f"This version of Yiding supports format version {FORMAT_VERSION}."
        )

    if validate:
        validate_project_data(project)

    return project


def apply_project_to_text(
    project: Mapping[str, Any],
    text: Any,
) -> Any:
    """
    Restore saved data onto an existing SegmentedText-compatible object.

    This intentionally does not call RawText/KanripoText constructors because
    KanripoText.__init__ fetches/parses its source again.
    """
    metadata = project["metadata"]
    source = project["source"]
    text_data = project["text"]
    workflow = project["workflow"]
    glossaries = project["glossaries"]

    text.original_title = metadata["original_title"]
    text.translated_title = metadata["translated_title"]
    text.translation_language = metadata["translation_language"]
    text.properties = deepcopy(metadata["properties"])
    text.is_structured = metadata["is_structured"]
    text.is_punctuated = metadata["is_punctuated"]
    text.is_translated = metadata["is_translated"]
    text.is_term_extracted = (metadata["is_term_extracted"])
    text.is_glossary_extracted = (metadata["is_glossary_extracted"])

    text.segments = deepcopy(text_data["segments"])
    text.unpunctuated_segments = deepcopy(
        text_data["unpunctuated_segments"]
    )
    text.punctuated_segments = deepcopy(
        text_data["punctuated_segments"]
    )
    text.translated_segments = deepcopy(
        text_data["translated_segments"]
    )
    text.page_labels = deepcopy(text_data["page_labels"])
    text.page_lines = deepcopy(text_data["page_lines"])
    text.segment_labels = deepcopy(text_data["segment_labels"])
    text.glosses = deepcopy(text_data["glosses"])
    text.punctuated_glosses = deepcopy(
        text_data["punctuated_glosses"]
    )

    text.settings = deepcopy(project["settings"])

    text.working_i = workflow["working_i"]
    text.working_text = workflow["working_text"]
    text.working_labels = deepcopy(workflow["working_labels"])
    text.translation_i = workflow["translation_i"]
    text.term_extraction_i = (workflow["term_extraction_i"])
    text.glossary_extraction_i = (workflow["glossary_extraction_i"])

    text.translation_glossary = glossary_rows_to_dict(
        glossaries["translation_glossary"]["rows"]
    )
    text.new_glossary = glossary_rows_to_dict(
        glossaries["new_glossary"]["rows"]
    )
    text.extracted_terms = glossary_rows_to_dict(glossaries["extracted_terms"]["rows"])
    text.extracted_glossary=glossary_rows_to_dict(glossaries["extracted_glossary"]["rows"])
    text.log = project["log"]

    if project["text_type"] == "raw":
        unpunctuated_source = source["unpunctuated"]
        punctuated_source = source["punctuated"]
        translation_source = source["translation"]
        text.unpunctuated_divider = (unpunctuated_source["divider"])
        text.unpunctuated_keep_divider = (unpunctuated_source["keep_divider"])
        text.unpunctuated_divider_is_whole_line = (
            unpunctuated_source["divider_is_whole_line"])
        text.punctuated_divider = (punctuated_source["divider"])
        text.punctuated_keep_divider = (punctuated_source["keep_divider"])
        text.punctuated_divider_is_whole_line = (punctuated_source["divider_is_whole_line"])
        text.translation_divider = (translation_source["divider"])
        text.translation_keep_divider = (translation_source["keep_divider"])
        text.translation_divider_is_whole_line = (
            translation_source["divider_is_whole_line"])
    
    elif project["text_type"] == "kanripo":
        text.kanripo_code = source["kanripo_code"]
        text.glosses_on = source["glosses_on"]

    # Keep project history/provenance available without forcing changes to
    # SegmentedText immediately.
    text.yiding_runs = deepcopy(project["runs"])
    text.yiding_provenance = deepcopy(project["provenance"])
    text.yiding_glossary_source = deepcopy(glossaries["source"])

    return text


def load_project_into_text(
    path: str | Path,
    text: Any,
    *,
    validate: bool = True,
) -> Any:
    project = load_project(
        path,
        validate=validate,
    )

    return apply_project_to_text(
        project,
        text,
    )


# ---------------------------------------------------------
# EXPORT
# ---------------------------------------------------------

def export_glossary_csv(
    glossary,
    path: str | Path,
) -> Path:

    path = Path(path)

    with path.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as file:

        writer = csv.writer(file)

        for chinese, value in glossary.items():

            pinyin = value[0]
            translations = value[1]

            writer.writerow([
                chinese,
                pinyin,
                *translations,
            ])

    return path


def export_log_txt(
    log: str,
    path: str | Path,
) -> Path:

    path = Path(path)

    path.write_text(
        log,
        encoding="utf-8",
    )

    return path


# ---------------------------------------------------------
# TEXT EXPORT
# ---------------------------------------------------------

def _format_export_settings(
    settings: Mapping[str, Any],
    indent: int = 0,
) -> list[str]:

    lines = []
    prefix = "    " * indent

    for key, value in settings.items():

        if isinstance(value, Mapping):

            lines.append(
                f"{prefix}{key}:"
            )

            lines.extend(
                _format_export_settings(
                    value,
                    indent + 1,
                )
            )

        else:

            if isinstance(value, bool):
                value = str(value).lower()

            elif value is None:
                value = "none"

            else:
                value = str(value)

            if "\n" in value:

                lines.append(
                    f"{prefix}{key}:"
                )

                for line in value.splitlines():

                    lines.append(
                        f"{prefix}    {line}"
                    )

            else:

                lines.append(
                    f"{prefix}{key}: {value}"
                )

    return lines


def _export_runs_for_text(
    project: Mapping[str, Any],
    text_kind: str,
) -> list[Mapping[str, Any]]:

    provenance_key = {
        "punctuated":
            "punctuated_segment_run_ids",

        "translation":
            "translated_segment_run_ids",
    }.get(text_kind)

    if provenance_key is None:
        return []

    run_ids = (
        project["provenance"]
        .get(provenance_key, [])
    )

    # Preserve the order in which the runs occur.
    unique_run_ids = []

    for run_id in run_ids:

        if (
            run_id is not None
            and run_id not in unique_run_ids
        ):
            unique_run_ids.append(run_id)

    runs_by_id = {
        run["id"]: run
        for run in project.get("runs", [])
    }

    return [
        runs_by_id[run_id]
        for run_id in unique_run_ids
        if run_id in runs_by_id
    ]


def export_text_txt(
    project: Mapping[str, Any],
    text_kind: str,
    path: str | Path,
) -> Path:

    path = Path(path)

    metadata = project["metadata"]
    text_data = project["text"]

    # SELECT TEXT

    if text_kind == "unpunctuated":

        segments = (
            text_data["unpunctuated_segments"]
            or text_data["segments"]
        )

        title = (
            metadata["original_title"]
            or "Untitled"
        )

        section_name = "Unpunctuated Text"

    elif text_kind == "punctuated":

        segments = (
            text_data["punctuated_segments"]
        )

        title = (
            metadata["original_title"]
            or "Untitled"
        )

        section_name = "Punctuated Text"

    elif text_kind == "translation":

        segments = (
            text_data["translated_segments"]
        )

        title = (
            metadata["translated_title"]
            or metadata["original_title"]
            or "Untitled"
        )

        section_name = "Translation"

    else:

        raise YidingFileError(
            f"Unknown TXT export type: {text_kind}"
        )

    if not segments:

        raise YidingFileError(
            f"No {section_name.lower()} exists "
            "to export."
        )

    # HEADER

    lines = [
        f"Title: {title}",
        "",
        "Properties:",
    ]

    properties = metadata.get(
        "properties",
        []
    )

    if properties:

        for prop in properties:
            lines.append(
                str(prop).rstrip()
            )

    else:

        lines.append("None")

    # SETTINGS USED TO PRODUCE THE TEXT

    runs = _export_runs_for_text(
        project,
        text_kind,
    )

    for number, run in enumerate(
        runs,
        start=1,
    ):

        operation_name = (
            run["operation"]
            .replace("_", " ")
            .title()
        )

        lines.append("")

        if len(runs) == 1:

            lines.append(
                f"{operation_name} settings:"
            )

        else:

            lines.append(
                f"{operation_name} settings "
                f"(run {number}):"
            )

        lines.extend(
            _format_export_settings(
                run["settings"],
                indent=1,
            )
        )

    # TEXT

    lines.extend([
        "",
        section_name + ":",
        "",
    ])

    labels = text_data.get(
        "segment_labels",
        []
    )

    # Only use labels if every exported segment
    # actually has a corresponding label.
    use_labels = (
        bool(labels)
        and len(labels) == len(segments)
    )

    for i, segment in enumerate(segments):

        if use_labels:

            lines.append(
                str(labels[i])
            )

        lines.append(
            str(segment)
        )

        lines.append("")

    path.write_text(
        "\n".join(lines).rstrip() + "\n",
        encoding="utf-8",
    )

    return path



# ---------------------------------------------------------
# DOCX EXPORT
# ---------------------------------------------------------

def _docx_export_runs(
    project: Mapping[str, Any],
    export_kind: str,
) -> list[Mapping[str, Any]]:

    text_kinds = []

    if export_kind in (
        "punctuated",
        "punctuated_translation",
    ):
        text_kinds.append("punctuated")

    if export_kind in (
        "translation",
        "punctuated_translation",
    ):
        text_kinds.append("translation")

    runs = []
    seen_ids = set()

    for text_kind in text_kinds:

        for run in _export_runs_for_text(
            project,
            text_kind,
        ):

            run_id = run["id"]

            if run_id not in seen_ids:
                runs.append(run)
                seen_ids.add(run_id)

    return runs


def _set_source_font(
    run,
    size=12,
):

    run.font.name = "PMingLiU"
    run.font.size = Pt(size)

    fonts = run._element.rPr.rFonts

    fonts.set(
        qn("w:eastAsia"),
        "PMingLiU",
    )


def _set_latin_font(
    run,
    size=12,
):

    run.font.name = "Times New Roman"
    run.font.size = Pt(size)

    fonts = run._element.rPr.rFonts

    fonts.set(
        qn("w:eastAsia"),
        "Times New Roman",
    )


def _add_docx_title_page(
    doc,
    metadata: Mapping[str, Any],
    export_kind: str,
) -> bool:

    original_title = metadata.get(
        "original_title"
    )

    translated_title = metadata.get(
        "translated_title"
    )

    include_translated_title = (
        export_kind in (
            "translation",
            "punctuated_translation",
        )
        and translated_title
    )

    if not original_title and not include_translated_title:
        return False

    if original_title:

        paragraph = doc.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            "\n\n\n" + original_title
        )

        _set_source_font(
            run,
            28,
        )

        run.font.bold = True
        run.font.color.rgb = RGBColor(
            0,
            0,
            0,
        )

    if include_translated_title:

        paragraph = doc.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            translated_title
        )

        _set_latin_font(
            run,
            18,
        )

        run.font.bold = True
        run.font.color.rgb = RGBColor(
            0,
            0,
            0,
        )

    doc.add_page_break()

    return True


def _add_docx_metadata_page(
    doc,
    project: Mapping[str, Any],
    export_kind: str,
) -> bool:

    metadata = project["metadata"]

    properties = metadata.get(
        "properties",
        []
    )

    runs = _docx_export_runs(
        project,
        export_kind,
    )

    translation_language = None

    if export_kind in (
        "translation",
        "punctuated_translation",
    ):
        translation_language = metadata.get(
            "translation_language"
        )

    if (
        not properties
        and not runs
        and not translation_language
    ):
        return False

    # -----------------------------------------------------
    # TEXT PROPERTIES
    # -----------------------------------------------------

    if properties:

        paragraph = doc.add_paragraph()

        heading = paragraph.add_run(
            "Text Properties"
        )

        _set_latin_font(
            heading,
            11,
        )

        heading.font.bold = True

        for prop in properties:

            paragraph = doc.add_paragraph()

            run = paragraph.add_run(
                str(prop).rstrip()
            )

            _set_latin_font(
                run,
                9,
            )

    # -----------------------------------------------------
    # TRANSLATION LANGUAGE
    # -----------------------------------------------------

    if translation_language:

        paragraph = doc.add_paragraph()

        heading = paragraph.add_run(
            "Translation Language: "
        )

        _set_latin_font(
            heading,
            9,
        )

        heading.font.bold = True

        value = paragraph.add_run(
            str(translation_language)
        )

        _set_latin_font(
            value,
            9,
        )

    # -----------------------------------------------------
    # TRANSFORMATION SETTINGS
    # -----------------------------------------------------

    for number, run_data in enumerate(
        runs,
        start=1,
    ):

        paragraph = doc.add_paragraph()

        operation_name = (
            run_data["operation"]
            .replace("_", " ")
            .title()
        )

        heading_text = (
            f"{operation_name} Settings"
        )

        if len(runs) > 1:
            heading_text += (
                f" — Run {number}"
            )

        heading = paragraph.add_run(
            heading_text
        )

        _set_latin_font(
            heading,
            11,
        )

        heading.font.bold = True

        settings_lines = (
            _format_export_settings(
                run_data["settings"]
            )
        )

        paragraph = doc.add_paragraph()

        settings_run = paragraph.add_run(
            "\n".join(settings_lines)
        )

        _set_latin_font(
            settings_run,
            8,
        )

    doc.add_page_break()

    return True


def _add_docx_label(
    container,
    label,
):

    paragraph = (
        container.add_paragraph()
    )

    run = paragraph.add_run(
        str(label)
    )

    _set_latin_font(
        run,
        10,
    )

    run.font.italic = True


def _add_docx_segment(
    container,
    segment,
    *,
    latin=False,
):

    paragraph = (
        container.add_paragraph()
    )

    run = paragraph.add_run(
        str(segment)
    )

    if latin:

        _set_latin_font(
            run,
            12,
        )

    else:

        _set_source_font(
            run,
            12,
        )


def export_text_docx(
    project: Mapping[str, Any],
    export_kind: str,
    path: str | Path,
    *,
    parallel_table: bool = True,
) -> Path:

    path = Path(path)

    metadata = project["metadata"]
    text_data = project["text"]

    unpunctuated = (
        text_data["unpunctuated_segments"]
        or text_data["segments"]
    )

    punctuated = (
        text_data["punctuated_segments"]
    )

    translated = (
        text_data["translated_segments"]
    )

    labels = text_data.get(
        "segment_labels",
        []
    )

    # ---------------------------------------------------------
    # VALIDATE EXPORT TYPE
    # ---------------------------------------------------------

    if export_kind == "unpunctuated":

        if not unpunctuated:
            raise YidingFileError(
                "No unpunctuated text exists "
                "to export."
            )

    elif export_kind == "punctuated":

        if not punctuated:
            raise YidingFileError(
                "No punctuated text exists "
                "to export."
            )

    elif export_kind == "translation":

        if not translated:
            raise YidingFileError(
                "No translation exists "
                "to export."
            )

    elif export_kind == (
        "punctuated_translation"
    ):

        if not punctuated or not translated:
            raise YidingFileError(
                "Both punctuated text and "
                "translation are required."
            )

        if len(punctuated) != len(translated):

            raise YidingFileError(
                "Punctuated text and translation "
                "have different numbers of segments."
            )

    else:

        raise YidingFileError(
            "Unknown DOCX export type: "
            f"{export_kind}"
        )

    # ---------------------------------------------------------
    # CREATE DOCUMENT
    # ---------------------------------------------------------

    doc = Document()

    normal_style = doc.styles["Normal"]

    normal_style.font.name = (
        "Times New Roman"
    )

    normal_style.font.size = Pt(12)

    # ---------------------------------------------------------
    # LANDSCAPE FOR PARALLEL TEXT
    # ---------------------------------------------------------

    if (
        export_kind == "punctuated_translation"
        and parallel_table
    ):

        section = doc.sections[0]

        section.orientation = (
            WD_ORIENTATION.LANDSCAPE
        )

        (
            section.page_width,
            section.page_height,
        ) = (
            section.page_height,
            section.page_width,
        )

    # ---------------------------------------------------------
    # TITLE
    # ---------------------------------------------------------

    _add_docx_title_page(
        doc,
        metadata,
        export_kind,
    )

    # ---------------------------------------------------------
    # PROPERTIES / SETTINGS
    # ---------------------------------------------------------

    _add_docx_metadata_page(
        doc,
        project,
        export_kind,
    )

    # ---------------------------------------------------------
    # PARALLEL PUNCTUATED + TRANSLATION
    # ---------------------------------------------------------

    if (
        export_kind == "punctuated_translation"
        and parallel_table
    ):

        table = doc.add_table(
            rows=len(translated),
            cols=2,
        )

        table.autofit = False

        section = doc.sections[0]

        usable_width = (
            section.page_width
            - section.left_margin
            - section.right_margin
        )

        source_width = (
            usable_width // 3
        )

        translation_width = (
            usable_width
            - source_width
        )

        table.columns[0].width = (
            source_width
        )

        table.columns[1].width = (
            translation_width
        )

        for row in table.rows:

            row.cells[0].width = (
                source_width
            )

            row.cells[1].width = (
                translation_width
            )

            row.allow_break_across_pages = (
                False
            )

        for i in range(
            len(translated)
        ):

            source_cell = table.cell(
                i,
                0,
            )

            translation_cell = table.cell(
                i,
                1,
            )

            if (
                i < len(labels)
                and labels[i]
            ):

                _add_docx_label(
                    source_cell,
                    labels[i],
                )

                _add_docx_label(
                    translation_cell,
                    labels[i],
                )

            _add_docx_segment(
                source_cell,
                punctuated[i],
            )

            _add_docx_segment(
                translation_cell,
                translated[i],
                latin=True,
            )

    elif export_kind == "punctuated_translation":

        for i in range(
            len(translated)
        ):

            if (
                i < len(labels)
                and labels[i]
            ):

                _add_docx_label(
                    doc,
                    labels[i],
                )

            _add_docx_segment(
                doc,
                punctuated[i],
            )

            _add_docx_segment(
                doc,
                translated[i],
                latin=True,
            )
        
    # ---------------------------------------------------------
    # SINGLE TEXT
    # ---------------------------------------------------------

    else:

        if export_kind == (
            "unpunctuated"
        ):

            segments = unpunctuated
            latin = False

        elif export_kind == (
            "punctuated"
        ):

            segments = punctuated
            latin = False

        else:

            segments = translated
            latin = True

        for i, segment in enumerate(
            segments
        ):

            if (
                i < len(labels)
                and labels[i]
            ):

                _add_docx_label(
                    doc,
                    labels[i],
                )

            _add_docx_segment(
                doc,
                segment,
                latin=latin,
            )

    # ---------------------------------------------------------
    # SAVE
    # ---------------------------------------------------------

    doc.save(
        str(path)
    )

    return path


