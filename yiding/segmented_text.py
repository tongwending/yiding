# ------------------------------------------------------------------------------------------
# segmented_text
# ------------------------------------------------------------------------------------------

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENTATION

# ------------------------------------------------------------------------------------------

CHARACTER_CAP = 120

# ------------------------------------------------------------------------------------------

class SegmentedText:
    
    def __init__(self):
        
        self.original_title = None
        self.translated_title = None
        self.properties = []
        self.settings = None
        
        self.segments = []
        self.unpunctuated_segments = []
        self.punctuated_segments = []
        self.translated_segments = []
        
        self.page_labels = []
        self.page_lines = []
        self.segment_labels = []
        
        self.glosses = []
        self.punctuated_glosses = []
        
        self.translation_glossary = {}

        self.working_i = 0
        self.working_text = ""
        self.working_labels = []
        self.translation_i = 0
        self.log = ""

        self.is_structured = False # True means it is divided into meaningful sections.
        self.is_punctuated = False      
                                    
    def full_title(self):
        fulltitle = (self.original_title if self.original_title else "Untitled")\
                  + (f" - {self.translated_title}" if self.translated_title else "")
        return fulltitle

    def empty_translation(self, empty_title = True):
        self.translation_i = 0
        self.translated_segments = []
        self.translation_glossary = {}
        if empty_title:
            self.translated_title = None

    def _update_log(self, text):
        self.log += f"{text}\n"; print(text)
            
# ------------------------------------------------------------------------------------------
# i/o behaviour
# ------------------------------------------------------------------------------------------

    def print_punctuation_settings(self):
        text = ("Punctuation settings:\n"
                f'facsimile_span: {self.settings["punctuation"]["facsimile_span"]}\n'
                f'punctuation_span: {self.settings["punctuation"]["punctuation_span"]}\n'
                f'max_unsegmented_span: {self.settings["punctuation"]["max_unsegmented_span"]}\n'
                f'max_punctuation_attempts: {self.settings["punctuation"]["max_punctuation_attempts"]}\n'
                f'cross_check: {self.settings["punctuation"]["cross_check"]}\n'
                f'punctuation: {self.settings["punctuation"]["punctuation"]}\n')
        if self.settings["punctuation"]["cross_check"]["enabled"]:
            text += (f'cross_examination: {self.settings["punctuation"]["cross_examination"]}\n'
                     f'cross_correction: {self.settings["punctuation"]["cross_correction"]}\n')
        text += f'guidelines:\n{self.settings["punctuation"]["guidelines"]}'
        return text

    def print_translation_settings(self):
        text = ("Translation settings:\n"
                f'language: {self.settings["translation"]["language"]}\n'
                f'glossary: {self.settings["translation"]["glossary"]}\n'
                f'translation_span: {self.settings["translation"]["translation_span"]}\n'
                f'cross_check: {self.settings["translation"]["cross_check"]}\n'
                f'llm_glossary_selection: {self.settings["translation"]["llm_glossary_selection"]}\n'
                f'translation: {self.settings["translation"]["translation"]}\n'
                f'glossary_extraction: {self.settings["translation"]["glossary_extraction"]}\n')
        if self.settings["translation"]["cross_check"]["enabled"]:
            text += (f'cross_examination: {self.settings["translation"]["cross_examination"]}\n'
                     f'cross_correction: {self.settings["translation"]["cross_correction"]}\n')
        if self.settings["translation"]["llm_glossary_selection"]["enabled"]:
            text += (f'glossary_selection: {self.settings["translation"]["glossary_selection"]}\n')

        text += f'guidelines:\n{self.settings["translation"]["guidelines"]}'
        return text

# ------------------------------------------------------------------------------------------

    def print_to_file(self, punctuation = True, translation = True,
                      output_file = "docx", table = True):
        if output_file == "docx":
            self.print_to_docx(punctuation = punctuation, translation = translation,
                               table = table)
   
# ------------------------------------------------------------------------------------------


    def print_to_docx(self, punctuation = True, translation = True, table = True):

        if punctuation and translation:
            if len(self.punctuated_segments) != len(self.translated_segments):
                raise ValueError("Error: Segment lists' length mismatch.")

        # Assign Word document class.
        doc = Document()

        if table:
            section = doc.sections[0]
            section.orientation = WD_ORIENTATION.LANDSCAPE
            section.page_width, section.page_height= section.page_height, section.page_width

        normal_style = doc.styles['Normal']
        normal_style.font.name = 'Times New Roman'
        normal_style.font.size = Pt(12)

        # first page:
        
        # 1) Chinese title
        if self.original_title and punctuation:
            p = doc.add_paragraph()   # create an empty paragraph.
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_cn = p.add_run("\n\n\n" + self.original_title)
            run_cn.font.name = 'PMingLiU'
            run_cn.font.size = Pt(28)
            run_cn.font.bold = True
            run_cn.font.color.rgb = RGBColor(0, 0, 0)
            r_cn = run_cn._element.rPr.rFonts
            r_cn.set(qn('w:eastAsia'), 'PMingLiU')

            if self.translated_title and translation:
                # 2) English title
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_en = p.add_run(self.translated_title)
                run_en.font.name = 'Times New Roman'
                run_en.font.size = Pt(18)
                run_en.font.bold = True
                run_en.font.color.rgb = RGBColor(0, 0, 0)
                r_en = run_en._element.rPr.rFonts
                r_en.set(qn('w:eastAsia'), 'Times New Roman')

            doc.add_page_break()

        # second page:

        # 3) Text properties and prompting settings (leaves empty page if there is nothing)
        p = doc.add_paragraph()
        properties_text = ""
        settings_text = ""
        if self.properties:
            properties_text += "Text properties:"
            for prop in self.properties:
                properties_text += f"\n{prop.rstrip()}"
        settings_text = f"{self.print_punctuation_settings()}"
        if translation:
            settings_text += f"\n{self.print_translation_settings()}"
        run_num = p.add_run(f"{properties_text}\n\n{settings_text}")
        run_num.font.name = 'Times New Roman'
        run_num.font.size = Pt(7)
        run_num.font.color.rgb = RGBColor(0, 0, 0)
        r_num = run_num._element.rPr.rFonts
        r_num.set(qn('w:eastAsia'), 'Times New Roman')
        
        doc.add_page_break()

        # third page and on: fill in the segments
        
        def insert_label(label):
            if punctuation and translation and table:
                p = cell.paragraphs[-1]
            else:
                p = doc.add_paragraph()
            run = p.add_run(f"{label}\n")
            run.font.size = Pt(12)
            run.font.name = 'Times New Roman'
            run.font.italic = True

        def insert_segment(segment, latin = False):
            if punctuation and translation and table:
                p = cell.paragraphs[-1]
            else:
                p = doc.add_paragraph()
            run = p.add_run(f"{segment}"
                          +("\n\n" if not table else ""))
            run.font.size = Pt(12)
            if latin:
                run.font.name = 'Times New Roman'
            else:
                run.font.name = 'PMingLiU'
                # Makes sure Word treats it as East Asian font
                r = run._element.rPr.rFonts
                r.set(qn('w:eastAsia'), 'PMingLiU')
                
        if punctuation and translation and table:
            t = doc.add_table(rows = len(self.translated_segments), cols = 2) # t for table

            # fix cell width ratio
            t.autofit = False
            usable_width = section.page_width - section.left_margin - section.right_margin
            col1_w = usable_width // 3
            col2_w = usable_width - col1_w  # = 2/3
            t.columns[0].width = col1_w
            t.columns[1].width = col2_w
            for row in t.rows:
                row.cells[0].width = col1_w
                row.cells[1].width = col2_w
                row.allow_break_across_pages = False

            for i in range(0, len(self.translated_segments)):
                # segement a cell
                cell = t.cell(i, 0)
                insert_label(self.segment_labels[i])
                insert_segment(self.punctuated_segments[i])
                # segment b cell
                cell = t.cell(i, 1)
                insert_label(self.segment_labels[i])
                insert_segment(self.translated_segments[i], latin = True)
        else:
            if punctuation:
                length = len(self.punctuated_segments)
            elif translation:
                length = len(self.translated_segments)
            else:
                length = 0
            for i in range(0, length):
                insert_label(self.segment_labels[i])
                if punctuation:
                    insert_segment(self.punctuated_segments[i])
                if translation:
                    insert_segment(self.translated_segments[i], latin = True)

        filetitle = (f"{self.full_title()[:CHARACTER_CAP]}"
                     + (" - Punctuated" if (punctuation and not translation) else "")
                     + (" - Translated" if (translation and not punctuation) else "")
                     + ".docx")
        doc.save(filetitle)   
        
# ------------------------------------------------------------------------------------------
