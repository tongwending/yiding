# ------------------------------------------------------------------------------------------
# segmented_text
# ------------------------------------------------------------------------------------------

from docx import Document
from docx.shared import Pt, RGBColor, Emu
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENTATION

# ------------------------------------------------------------------------------------------

class SegmentedText:
    
    def __init__(self):
        
        self.original_title = None
        self.translated_title = None
        self.properties = []
        
        self.segments = []
        self.unpunctuated_segments = []
        self.punctuated_segments = []
        self.translated_segments = []
        
        self.page_labels = []
        self.page_lines = []
        self.segment_labels = []
        
        self.glosses = []
        self.punctuated_glosses = []
        
        self.translation_glossary = {}

        self.working_i = 0
        self.working_text = ""
        self.working_labels = []
        self.translation_i = 0

        self.is_punctuated = False
        self.prompting_settings = {'punctuation': "", 'translation': ""}
        
                                    
    def full_title(self):
        fulltitle = (self.original_title if self.original_title else "Untitled")\
                  + (f" - {self.translated_title}" if self.translated_title else "")
        return fulltitle

    def empty_translation(self, empty_title = True):
        self.translation_i = 0
        self.translated_segments = []
        self.translation_glossary = {}
        if empty_title:
            self.translated_title = None
        self.prompting_settings['translation'] = ""
            
# ------------------------------------------------------------------------------------------
# i/o behaviour
# ------------------------------------------------------------------------------------------

    def print_punctuation_settings(self):
        text = (
            "Punctuation settings:\n"
            f"punctuation cross-check: {self.settings.PUNCTUATION_CROSS_CHECK}\n"
            f"fascimile span: {self.settings.FASCIMILE_SPAN}\n"
            f"punctuation span: {self.settings.PUNCTUATION_SPAN}\n"
            f"maximum unsegmented span: {self.settings.MAX_UNSEGMENTED_SPAN}\n"
            f"maximum punctuation attempts: {self.settings.MAX_PUNCTUATION_ATTEMPTS}\n"
            f"punctuation guidelines:\n{self.settings.PUNCTUATION_GUIDELINES}\n"
            "punctuation: "
                f"model: {self.settings.PUNCTUATION_MODEL}, "
                f"reasoning: {self.settings.PUNCTUATION_REASONING}, "
                f"verbosity: {self.settings.PUNCTUATION_VERBOSITY}, "
                f"temperature: {self.settings.PUNCTUATION_TEMPERATURE}, "
                f"top_p: {self.settings.PUNCTUATION_TOP_P}\n"
            "punctuation examination: "
                f"model: {self.settings.PUNCTUATION_EXAMINATION_MODEL}, "
                f"reasoning: {self.settings.PUNCTUATION_EXAMINATION_REASONING}, "
                f"verbosity: {self.settings.PUNCTUATION_EXAMINATION_VERBOSITY}, "
                f"temperature: {self.settings.PUNCTUATION_EXAMINATION_TEMPERATURE}, "
                f"top_p: {self.settings.PUNCTUATION_EXAMINATION_TOP_P}\n"
            "punctuation correction: "
                f"model: {self.settings.PUNCTUATION_CORRECTION_MODEL}, "
                f"reasoning: {self.settings.PUNCTUATION_CORRECTION_REASONING}, "
                f"verbosity: {self.settings.PUNCTUATION_CORRECTION_VERBOSITY}, "
                f"temperature: {self.settings.PUNCTUATION_CORRECTION_TEMPERATURE}, "
                f"top_p: {self.settings.PUNCTUATION_CORRECTION_TOP_P}")    
        return text

    def print_translation_settings(self):
        text = (
            "Translation settings:\n"
            +f"translation cross-check: {self.settings.TRANSLATION_CROSS_CHECK}\n"
            +(f"LLM-powered glossary selection: {self.settings.LLM_GLOSSARY_SELECTION}\n"
                if self.settings.LLM_GLOSSARY_SELECTION else "")
            +f"translation language: {self.settings.LANGUAGE}\n"
            +f"translation span: {self.settings.TRANSLATION_SPAN}\n"
            +f"glossary file: {self.settings.GLOSSARY_FILE}\n"
            +f"translation guidelines:\n{self.settings.TRANSLATION_GUIDELINES}\n"
            +(("glossary selection: "
                f"model: {self.settings.GLOSSARY_SELECTION_MODEL}, "
                f"reasoning: {self.settings.GLOSSARY_SELECTION_REASONING}, "
                f"verbosity: {self.settings.GLOSSARY_SELECTION_VERBOSITY}, "
                f"temperature: {self.settings.GLOSSARY_SELECTION_TEMPERATURE}, "
                f"top_p: {self.settings.GLOSSARY_SELECTION_TOP_P}\n")
                if self.settings.LLM_GLOSSARY_SELECTION else "")
            +("translation: "
                f"model: {self.settings.TRANSLATION_MODEL}, "
                f"reasoning: {self.settings.TRANSLATION_REASONING}, "
                f"verbosity: {self.settings.TRANSLATION_VERBOSITY}, "
                f"temperature: {self.settings.TRANSLATION_TEMPERATURE}, "
                f"top_p: {self.settings.TRANSLATION_TOP_P}\n")
            +("translation examination: "
                f"model: {self.settings.TRANSLATION_EXAMINATION_MODEL}, "
                f"reasoning: {self.settings.TRANSLATION_EXAMINATION_REASONING}, "
                f"verbosity: {self.settings.TRANSLATION_EXAMINATION_VERBOSITY}, "
                f"temperature: {self.settings.TRANSLATION_EXAMINATION_TEMPERATURE}, "
                f"top_p: {self.settings.TRANSLATION_EXAMINATION_TOP_P}\n")
            +("translation correction: "
                f"model: {self.settings.TRANSLATION_CORRECTION_MODEL}, "
                f"reasoning: {self.settings.TRANSLATION_CORRECTION_REASONING}, "
                f"verbosity: {self.settings.TRANSLATION_CORRECTION_VERBOSITY}, "
                f"temperature: {self.settings.TRANSLATION_CORRECTION_TEMPERATURE}, "
                f"top_p: {self.settings.TRANSLATION_CORRECTION_TOP_P}\n")
            +("glossary extraction: "
                f"model: {self.settings.GLOSSARY_EXTRACTION_MODEL}, "
                f"reasoning: {self.settings.GLOSSARY_EXTRACTION_REASONING}, "
                f"verbosity: {self.settings.GLOSSARY_EXTRACTION_VERBOSITY}, "
                f"temperature: {self.settings.GLOSSARY_EXTRACTION_TEMPERATURE}, "
                f"top_p: {self.settings.GLOSSARY_EXTRACTION_TOP_P}")
            )
        return text

# ------------------------------------------------------------------------------------------

    def print_to_file(self, punctuation = True, translation = True,
                      output_file = "docx", table = True):
        if output_file == "docx":
            self.print_to_docx(punctuation = punctuation, translation = translation,
                               table = table)
   
# ------------------------------------------------------------------------------------------


    def print_to_docx(self, punctuation = True, translation = True, table = True):

        if punctuation and translation:
            if len(self.punctuated_segments) != len(self.translated_segments):
                raise ValueError("Error: Segment lists' length mismatch.")

        # Assign Word document class.
        doc = Document()

        if table:
            section = doc.sections[0]
            section.orientation = WD_ORIENTATION.LANDSCAPE
            section.page_width, section.page_height= section.page_height, section.page_width

        normal_style = doc.styles['Normal']
        normal_style.font.name = 'Times New Roman'
        normal_style.font.size = Pt(12)

        # first page:
        
        # 1) Chinese title
        if self.original_title and punctuation:
            p = doc.add_paragraph()   # create an empty paragraph.
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_cn = p.add_run("\n\n\n" + self.original_title)
            run_cn.font.name = 'PMingLiU'
            run_cn.font.size = Pt(28)
            run_cn.font.bold = True
            run_cn.font.color.rgb = RGBColor(0, 0, 0)
            r_cn = run_cn._element.rPr.rFonts
            r_cn.set(qn('w:eastAsia'), 'PMingLiU')

            if self.translated_title and translation:
                # 2) English title
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_en = p.add_run(self.translated_title)
                run_en.font.name = 'Times New Roman'
                run_en.font.size = Pt(18)
                run_en.font.bold = True
                run_en.font.color.rgb = RGBColor(0, 0, 0)
                r_en = run_en._element.rPr.rFonts
                r_en.set(qn('w:eastAsia'), 'Times New Roman')

            doc.add_page_break()

        # second page:

        # 3) Text properties and prompting settings (leaves empty page if there is nothing)
        p = doc.add_paragraph()
        properties_text = ""
        settings_text = ""
        if self.properties:
            properties_text += "Text properties:"
            for prop in self.properties:
                properties_text += f"\n{prop.rstrip()}"
        settings_text = f"{self.print_punctuation_settings()}"
        if translation:
            settings_text += f"\n\n{self.print_translation_settings()}"
        run_num = p.add_run(f"{properties_text}\n\n{settings_text}")
        run_num.font.name = 'Times New Roman'
        run_num.font.size = Pt(6)
        run_num.font.color.rgb = RGBColor(0, 0, 0)
        r_num = run_num._element.rPr.rFonts
        r_num.set(qn('w:eastAsia'), 'Times New Roman')
        
        doc.add_page_break()

        # third page and on: fill in the segments
        
        def insert_label(label):
            if punctuation and translation and table:
                p = cell.paragraphs[-1]
            else:
                p = doc.add_paragraph()
            run = p.add_run(f"{label}\n")
            run.font.size = Pt(12)
            run.font.name = 'Times New Roman'
            run.font.italic = True

        def insert_segment(segment, latin = False):
            if punctuation and translation and table:
                p = cell.paragraphs[-1]
            else:
                p = doc.add_paragraph()
            run = p.add_run(f"{segment}"
                          +("\n\n" if not table else ""))
            run.font.size = Pt(12)
            if latin:
                run.font.name = 'Times New Roman'
            else:
                run.font.name = 'PMingLiU'
                # Makes sure Word treats it as East Asian font
                r = run._element.rPr.rFonts
                r.set(qn('w:eastAsia'), 'PMingLiU')
                
        if punctuation and translation and table:
            t = doc.add_table(rows = len(self.translated_segments), cols = 2) # t for table

            # fix cell width ratio
            t.autofit = False
            usable_width = section.page_width - section.left_margin - section.right_margin
            col1_w = usable_width // 3
            col2_w = usable_width - col1_w  # = 2/3
            t.columns[0].width = col1_w
            t.columns[1].width = col2_w
            for row in t.rows:
                row.cells[0].width = col1_w
                row.cells[1].width = col2_w
                row.allow_break_across_pages = False

            for i in range(0, len(self.translated_segments)):
                # segement a cell
                cell = t.cell(i, 0)
                insert_label(self.segment_labels[i])
                insert_segment(self.punctuated_segments[i])
                # segment b cell
                cell = t.cell(i, 1)
                insert_label(self.segment_labels[i])
                insert_segment(self.translated_segments[i], latin = True)
        else:
            if punctuation:
                length = len(self.punctuated_segments)
            elif translation:
                length = len(self.translated_segments)
            else:
                length = 0
            for i in range(0, length):
                insert_label(self.segment_labels[i])
                if punctuation:
                    insert_segment(self.punctuated_segments[i])
                if translation:
                    insert_segment(self.translated_segments[i], latin = True)

        filetitle = (f"{self.full_title()}"
                     + (" - Punctuated" if (punctuation and not translation) else "")
                     + (" - Translated" if (translation and not punctuation) else "")
                     + ".docx")
        doc.save(filetitle)   
        
# ------------------------------------------------------------------------------------------
